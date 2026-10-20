# %% Exportación del informe (Markdown → DOCX + PDF)
"""Genera informe_examen_final.docx y .pdf a partir de plantilla_examen_final.md."""

from __future__ import annotations

import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

RUTA_DOCS = Path(__file__).resolve().parent
MD_PATH = RUTA_DOCS / "plantilla_examen_final.md"
DOCX_PATH = RUTA_DOCS / "informe_examen_final.docx"
PDF_PATH = RUTA_DOCS / "informe_examen_final.pdf"


def _limpiar_md_inline(texto: str) -> str:
    texto = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", texto)
    texto = texto.replace("**", "").replace("__", "")
    texto = texto.replace("*", "")
    texto = re.sub(r"`([^`]+)`", r"\1", texto)
    texto = re.sub(r"\{\#[^}]+\}", "", texto)
    texto = re.sub(r"\$([^$]+)\$", r"\1", texto)
    return texto.strip()


def _es_encabezado(linea: str) -> tuple[int, str] | None:
    m = re.match(r"^(#{1,6})\s+(.*)$", linea.strip())
    if not m:
        return None
    nivel = len(m.group(1))
    texto = _limpiar_md_inline(m.group(2))
    # La plantilla usa "1. # **Capítulo..." — normalizar
    texto = re.sub(r"^\d+\.\s*", "", texto)
    return nivel, texto


def _set_run_font(run, size_pt: float = 12, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size_pt)
    run.bold = bold


def _aplicar_estilo_parrafo(parrafo, size_pt: float = 12, bold: bool = False, space_after: int = 8) -> None:
    fmt = parrafo.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_after = Pt(space_after)
    for run in parrafo.runs:
        _set_run_font(run, size_pt=size_pt, bold=bold)


def markdown_a_docx(md_text: str, destino: Path) -> None:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Times New Roman"
    estilo.font.size = Pt(12)
    estilo._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    estilo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    lineas = md_text.splitlines()
    i = 0
    en_tabla = False
    buffer_tabla: list[str] = []

    def volcar_tabla() -> None:
        nonlocal buffer_tabla, en_tabla
        if not buffer_tabla:
            en_tabla = False
            return
        filas = []
        for fila in buffer_tabla:
            if re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?$", fila.strip()):
                continue
            celdas = [c.strip() for c in fila.strip().strip("|").split("|")]
            filas.append([_limpiar_md_inline(c) for c in celdas])
        if filas:
            tabla = doc.add_table(rows=len(filas), cols=len(filas[0]))
            tabla.style = "Table Grid"
            for r_idx, fila in enumerate(filas):
                for c_idx, celda in enumerate(fila):
                    cell = tabla.rows[r_idx].cells[c_idx]
                    cell.text = celda
                    for p in cell.paragraphs:
                        for run in p.runs:
                            _set_run_font(run, size_pt=11, bold=(r_idx == 0))
            doc.add_paragraph("")
        buffer_tabla = []
        en_tabla = False

    while i < len(lineas):
        cruda = lineas[i]
        linea = cruda.rstrip()

        if linea.strip().startswith("|"):
            en_tabla = True
            buffer_tabla.append(linea)
            i += 1
            continue
        if en_tabla:
            volcar_tabla()

        if not linea.strip() or linea.strip() == "#":
            i += 1
            continue

        # Encabezados tipo "1. # **Capítulo" o "## título"
        if re.match(r"^\d+\.\s+#+\s+", linea.strip()) or linea.strip().startswith("#"):
            # Unificar "N. # texto" → encabezado
            normalizada = re.sub(r"^\d+\.\s+", "", linea.strip())
            info = _es_encabezado(normalizada if normalizada.startswith("#") else "# " + normalizada)
            if info:
                nivel, texto = info
                if nivel <= 1:
                    p = doc.add_heading(texto, level=1)
                elif nivel == 2:
                    p = doc.add_heading(texto, level=2)
                else:
                    p = doc.add_heading(texto, level=min(nivel, 3))
                for run in p.runs:
                    _set_run_font(run, size_pt=14 if nivel <= 1 else 12, bold=True)
                i += 1
                continue

        # Lista con viñetas
        if re.match(r"^(\*|-|\d+\.)\s+", linea.strip()):
            texto = re.sub(r"^(\*|-|\d+\.)\s+", "", linea.strip())
            texto = _limpiar_md_inline(texto)
            p = doc.add_paragraph(texto, style="List Bullet")
            _aplicar_estilo_parrafo(p, size_pt=12)
            i += 1
            continue

        # Fórmulas en bloque \[ ... \]
        if "\\[" in linea or linea.strip().startswith("$$"):
            bloque = [linea]
            i += 1
            while i < len(lineas) and ("\\]" not in lineas[i] and "$$" not in lineas[i]):
                bloque.append(lineas[i])
                i += 1
            if i < len(lineas):
                bloque.append(lineas[i])
                i += 1
            texto = _limpiar_md_inline(" ".join(bloque).replace("\\[", "").replace("\\]", "").replace("$$", ""))
            p = doc.add_paragraph(texto)
            _aplicar_estilo_parrafo(p, size_pt=11)
            continue

        # Párrafo normal; líneas en negrita sueltas del front-matter
        texto = _limpiar_md_inline(linea)
        if not texto:
            i += 1
            continue
        es_titulo_portada = linea.strip().startswith("**") and linea.strip().endswith("**")
        p = doc.add_paragraph(texto)
        _aplicar_estilo_parrafo(p, size_pt=14 if es_titulo_portada else 12, bold=es_titulo_portada)
        i += 1

    if en_tabla:
        volcar_tabla()

    doc.save(destino)


def _md_a_flowables(md_text: str):
    styles = getSampleStyleSheet()
    estilo_normal = ParagraphStyle(
        "InformeNormal",
        parent=styles["Normal"],
        fontName="Times-Roman",
        fontSize=11,
        leading=16.5,
        spaceAfter=8,
    )
    estilo_h1 = ParagraphStyle(
        "InformeH1",
        parent=styles["Heading1"],
        fontName="Times-Bold",
        fontSize=14,
        leading=18,
        spaceBefore=14,
        spaceAfter=8,
    )
    estilo_h2 = ParagraphStyle(
        "InformeH2",
        parent=styles["Heading2"],
        fontName="Times-Bold",
        fontSize=12,
        leading=16,
        spaceBefore=10,
        spaceAfter=6,
    )
    estilo_h3 = ParagraphStyle(
        "InformeH3",
        parent=styles["Heading3"],
        fontName="Times-Bold",
        fontSize=11,
        leading=15,
        spaceBefore=8,
        spaceAfter=4,
    )
    estilo_bullet = ParagraphStyle(
        "InformeBullet",
        parent=estilo_normal,
        leftIndent=18,
        bulletIndent=6,
    )

    story = []
    for cruda in md_text.splitlines():
        linea = cruda.rstrip()
        if not linea.strip() or linea.strip() == "#":
            story.append(Spacer(1, 4))
            continue

        if linea.strip().startswith("|"):
            texto = html.escape(_limpiar_md_inline(linea.replace("|", " | ")))
            story.append(Paragraph(texto, estilo_normal))
            continue

        if re.match(r"^\d+\.\s+#+\s+", linea.strip()) or linea.strip().startswith("#"):
            normalizada = re.sub(r"^\d+\.\s+", "", linea.strip())
            info = _es_encabezado(normalizada if normalizada.startswith("#") else "# " + normalizada)
            if info:
                nivel, texto = info
                estilo = estilo_h1 if nivel <= 1 else estilo_h2 if nivel == 2 else estilo_h3
                story.append(Paragraph(html.escape(texto), estilo))
                continue

        if re.match(r"^(\*|-|\d+\.)\s+", linea.strip()):
            texto = _limpiar_md_inline(re.sub(r"^(\*|-|\d+\.)\s+", "", linea.strip()))
            story.append(Paragraph("• " + html.escape(texto), estilo_bullet))
            continue

        texto = _limpiar_md_inline(linea)
        if texto:
            story.append(Paragraph(html.escape(texto), estilo_normal))
    return story


def markdown_a_pdf(md_text: str, destino: Path) -> None:
    doc = SimpleDocTemplate(
        str(destino),
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.5 * cm,
        bottomMargin=2.5 * cm,
        title="Informe Examen Final — Diagnóstico predictivo de bombas centrífugas",
    )
    doc.build(_md_a_flowables(md_text))


def main() -> None:
    md_text = MD_PATH.read_text(encoding="utf-8")
    markdown_a_docx(md_text, DOCX_PATH)
    markdown_a_pdf(md_text, PDF_PATH)
    print(f"DOCX: {DOCX_PATH} ({DOCX_PATH.stat().st_size} bytes)")
    print(f"PDF:  {PDF_PATH} ({PDF_PATH.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
